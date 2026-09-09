from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from uuid import UUID, uuid4

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.database import get_pool

MAX_ROWS_PER_SHEET = 5000
MAX_ERRORS = 200

SHEETS = {
    "Campuses": ["code", "name", "address", "phone", "email", "is_main"],
    "Academic Years": ["name", "start_date", "end_date", "is_current"],
    "Academic Terms": ["academic_year", "name", "term_number", "start_date", "end_date", "is_current"],
    "Departments": ["code", "name", "description"],
    "Class Levels": ["code", "name", "level_order"],
    "Streams": ["class_level", "code", "name", "capacity"],
    "Subjects": ["department", "code", "name", "short_name", "subject_type"],
    "Class Subjects": ["class_level", "subject", "is_compulsory"],
    "School Settings": ["setting_key", "setting_value", "value_type"],
}

DISPLAY_HEADERS = {
    "code": "code (REQUIRED)", "name": "name (REQUIRED)", "start_date": "start_date (REQUIRED)",
    "end_date": "end_date (REQUIRED)", "academic_year": "academic_year (REQUIRED)",
    "term_number": "term_number (REQUIRED)", "class_level": "class_level (REQUIRED)",
    "subject": "subject (REQUIRED)", "level_order": "level_order", "address": "address",
    "phone": "phone", "email": "email", "is_main": "is_main", "is_current": "is_current",
    "description": "description", "capacity": "capacity", "department": "department",
    "short_name": "short_name", "subject_type": "subject_type (REQUIRED)",
    "is_compulsory": "is_compulsory", "setting_key": "setting_key (REQUIRED)",
    "setting_value": "setting_value", "value_type": "value_type (REQUIRED)",
}


def _text(value, required=False):
    if value is None:
        if required:
            raise ValueError("value is required")
        return None
    value = str(value).strip()
    if required and not value:
        raise ValueError("value is required")
    return value or None


def _bool(value, default=False):
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    value = str(value).strip().lower()
    if value in {"true", "1", "yes", "y"}:
        return True
    if value in {"false", "0", "no", "n"}:
        return False
    raise ValueError("must be TRUE or FALSE")


def _int(value, default=None, minimum=None, maximum=None):
    if value is None or value == "":
        return default
    try:
        result = int(value)
    except (TypeError, ValueError):
        raise ValueError("must be an integer")
    if minimum is not None and result < minimum:
        raise ValueError(f"must be at least {minimum}")
    if maximum is not None and result > maximum:
        raise ValueError(f"must be at most {maximum}")
    return result


def _date(value):
    if value is None or value == "":
        raise ValueError("date is required")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError:
        raise ValueError("must be a date in YYYY-MM-DD format")


def _header_key(value):
    return str(value or "").strip().lower().replace(" (required)", "")


def _read_sheet(ws, expected):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_header_key(x) for x in rows[0]]
    if headers != expected:
        raise ValueError(f"Invalid columns. Expected: {', '.join(DISPLAY_HEADERS.get(x, x) for x in expected)}")
    if len(rows) - 1 > MAX_ROWS_PER_SHEET:
        raise ValueError(f"Maximum {MAX_ROWS_PER_SHEET} data rows allowed")
    result = []
    for row_number, values in enumerate(rows[1:], 2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        result.append((row_number, {expected[i]: values[i] if i < len(values) else None for i in range(len(expected))}))
    return result


def _parse_rows(workbook):
    errors = []
    parsed = {}
    for sheet, expected in SHEETS.items():
        if sheet not in workbook.sheetnames:
            errors.append({"sheet": sheet, "row": 1, "field": "sheet", "message": "Required sheet is missing"})
            continue
        try:
            parsed[sheet] = _read_sheet(workbook[sheet], expected)
        except ValueError as exc:
            errors.append({"sheet": sheet, "row": 1, "field": "columns", "message": str(exc)})
    return parsed, errors


def _parse_entity(sheet, row_no, row, errors):
    try:
        if sheet == "Campuses":
            return {"code": _text(row["code"], True), "name": _text(row["name"], True), "address": _text(row["address"]), "phone": _text(row["phone"]), "email": _text(row["email"]), "is_main": _bool(row["is_main"])}
        if sheet == "Academic Years":
            start, end = _date(row["start_date"]), _date(row["end_date"])
            if end <= start: raise ValueError("end_date must be after start_date")
            return {"name": _text(row["name"], True), "start_date": start, "end_date": end, "is_current": _bool(row["is_current"])}
        if sheet == "Academic Terms":
            start, end = _date(row["start_date"]), _date(row["end_date"])
            if end <= start: raise ValueError("end_date must be after start_date")
            return {"academic_year": _text(row["academic_year"], True), "name": _text(row["name"], True), "term_number": _int(row["term_number"], minimum=1, maximum=4), "start_date": start, "end_date": end, "is_current": _bool(row["is_current"])}
        if sheet == "Departments":
            return {"code": _text(row["code"], True), "name": _text(row["name"], True), "description": _text(row["description"])}
        if sheet == "Class Levels":
            return {"code": _text(row["code"], True), "name": _text(row["name"], True), "level_order": _int(row["level_order"], default=0)}
        if sheet == "Streams":
            return {"class_level": _text(row["class_level"], True), "code": _text(row["code"], True), "name": _text(row["name"], True), "capacity": _int(row["capacity"], minimum=1)}
        if sheet == "Subjects":
            subject_type = _text(row["subject_type"], True)
            if subject_type not in {"core", "elective", "co_curricular"}: raise ValueError("subject_type must be core, elective, or co_curricular")
            return {"department": _text(row["department"]), "code": _text(row["code"], True), "name": _text(row["name"], True), "short_name": _text(row["short_name"]), "subject_type": subject_type}
        if sheet == "Class Subjects":
            return {"class_level": _text(row["class_level"], True), "subject": _text(row["subject"], True), "is_compulsory": _bool(row["is_compulsory"], True)}
        if sheet == "School Settings":
            value_type = _text(row["value_type"], True)
            if value_type not in {"string", "integer", "boolean", "json"}: raise ValueError("value_type must be string, integer, boolean, or json")
            return {"setting_key": _text(row["setting_key"], True), "setting_value": _text(row["setting_value"]), "value_type": value_type}
    except ValueError as exc:
        errors.append({"sheet": sheet, "row": row_no, "field": "data", "message": str(exc)})
        return None
    return None


def _safe_excel(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _style_sheet(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in col) + 2, 14), 42)


async def _fetch_rows(tenant_id: UUID):
    pool = get_pool()
    data = {}
    queries = {
        "Campuses": ("SELECT code,name,address,phone,email,is_main FROM campuses WHERE tenant_id=%s ORDER BY name",),
        "Academic Years": ("SELECT name,start_date,end_date,is_current FROM academic_years WHERE tenant_id=%s ORDER BY start_date",),
        "Academic Terms": ("SELECT ay.name,atx.name,atx.term_number,atx.start_date,atx.end_date,atx.is_current FROM academic_terms atx JOIN academic_years ay ON ay.id=atx.academic_year_id WHERE atx.tenant_id=%s ORDER BY ay.start_date,atx.term_number",),
        "Departments": ("SELECT code,name,description FROM departments WHERE tenant_id=%s ORDER BY name",),
        "Class Levels": ("SELECT code,name,level_order FROM class_levels WHERE tenant_id=%s ORDER BY level_order,name",),
        "Streams": ("SELECT cl.code,s.code,s.name,s.capacity FROM streams s JOIN class_levels cl ON cl.id=s.class_level_id WHERE s.tenant_id=%s ORDER BY cl.level_order,s.name",),
        "Subjects": ("SELECT d.code,s.code,s.name,s.short_name,s.subject_type FROM subjects s LEFT JOIN departments d ON d.id=s.department_id WHERE s.tenant_id=%s ORDER BY s.name",),
        "Class Subjects": ("SELECT cl.code,s.code,cs.is_compulsory FROM class_subjects cs JOIN class_levels cl ON cl.id=cs.class_level_id JOIN subjects s ON s.id=cs.subject_id WHERE cs.tenant_id=%s ORDER BY cl.level_order,s.name",),
        "School Settings": ("SELECT setting_key,setting_value,value_type FROM school_settings WHERE tenant_id=%s ORDER BY setting_key",),
    }
    async with pool.acquire() as conn:
        async with conn.cursor() as cur:
            for sheet, (sql,) in queries.items():
                await cur.execute(sql, (str(tenant_id),))
                data[sheet] = await cur.fetchall()
    return data


def build_template():
    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("START HERE")
    readme.append(["ShuleLink School Structure Import/Export", ""])
    readme.append(["Import order", "Campuses → Academic Years → Terms → Departments → Class Levels → Streams → Subjects → Class Subjects → School Settings"])
    readme.append(["Rules", "Do not rename sheets or columns. Dates use YYYY-MM-DD. TRUE/FALSE are accepted. Parent references use codes/names shown in the workbook."])
    readme.append(["Modes", "Create New rejects duplicates. Upsert creates missing records and updates matching natural keys. No rows are deleted by import."])
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 100
    for sheet, fields in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append([DISPLAY_HEADERS[x] for x in fields])
        _style_sheet(ws)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def build_export(data):
    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("START HERE")
    readme.append(["ShuleLink School Structure Export", ""])
    readme.append(["Import-ready", "This export uses stable school-level natural keys rather than database UUIDs so it can be edited and re-imported safely."])
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 100
    for sheet, fields in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append([DISPLAY_HEADERS[x] for x in fields])
        for row in data.get(sheet, []):
            ws.append([_safe_excel(row[i]) for i in range(len(fields))])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in col) + 2, 14), 42)
        for row in ws.iter_rows():
            for cell in row:
                if cell.column in {2, 3, 4, 5} and "date" in str(ws.cell(1, cell.column).value).lower():
                    cell.number_format = "yyyy-mm-dd"
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def import_workbook(tenant_id: UUID, content: bytes, mode: str):
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "Import file is too large. Maximum size is 10 MB.")
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "The uploaded file is not a valid Excel workbook.")
    parsed, errors = _parse_rows(wb)
    wb.close()
    for sheet, rows in parsed.items():
        for row_no, raw in rows:
            _parse_entity(sheet, row_no, raw, errors)
            if len(errors) >= MAX_ERRORS:
                break
        if len(errors) >= MAX_ERRORS:
            break
    if errors:
        return {"valid": False, "mode": mode, "errors": errors[:MAX_ERRORS], "error_count": len(errors), "created": 0, "updated": 0}
    pool = get_pool()
    created = updated = 0
    async with pool.acquire() as conn:
        await conn.begin()
        try:
            async with conn.cursor() as cur:
                def sql_where(table, where):
                    return f"SELECT id FROM {table} WHERE tenant_id=%s AND {where} LIMIT 1"

                ids = {}
                async def find(table, where, args):
                    await cur.execute(sql_where(table, where), (str(tenant_id), *args))
                    row = await cur.fetchone()
                    return str(row[0]) if row else None

                async def require(table, where, args, label):
                    value = await find(table, where, args)
                    if not value:
                        raise ValueError(f"{label} not found")
                    return value

                for sheet in ["Campuses", "Academic Years", "Academic Terms", "Departments", "Class Levels", "Streams", "Subjects", "Class Subjects", "School Settings"]:
                    for row_no, raw in parsed.get(sheet, []):
                        item = _parse_entity(sheet, row_no, raw, [])
                        if sheet == "Campuses":
                            existing = await find("campuses", "code=%s", (item["code"],))
                            values = (item["name"], item["address"], item["phone"], item["email"], item["is_main"])
                            if existing and mode == "create": raise ValueError(f"Campuses row {row_no}: code already exists")
                            if existing: await cur.execute("UPDATE campuses SET name=%s,address=%s,phone=%s,email=%s,is_main=%s WHERE id=%s AND tenant_id=%s", (*values, existing, str(tenant_id))); updated += 1
                            else: await cur.execute("INSERT INTO campuses (id,tenant_id,code,name,address,phone,email,is_main) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)", (str(uuid4()),str(tenant_id),item["code"],*values)); created += 1
                        elif sheet == "Academic Years":
                            existing = await find("academic_years", "name=%s", (item["name"],))
                            if item["is_current"]: await cur.execute("UPDATE academic_years SET is_current=0 WHERE tenant_id=%s", (str(tenant_id),))
                            if existing and mode == "create": raise ValueError(f"Academic Years row {row_no}: name already exists")
                            values=(item["start_date"],item["end_date"],item["is_current"])
                            if existing: await cur.execute("UPDATE academic_years SET start_date=%s,end_date=%s,is_current=%s WHERE id=%s AND tenant_id=%s", (*values,existing,str(tenant_id))); updated+=1
                            else: await cur.execute("INSERT INTO academic_years (id,tenant_id,name,start_date,end_date,is_current) VALUES (%s,%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),item["name"],*values)); created+=1
                        elif sheet == "Academic Terms":
                            year_id = await require("academic_years", "name=%s", (item["academic_year"],), f"Academic year '{item['academic_year']}'")
                            existing = await find("academic_terms", "academic_year_id=%s AND term_number=%s", (year_id,item["term_number"]))
                            if item["is_current"]: await cur.execute("UPDATE academic_terms SET is_current=0 WHERE tenant_id=%s", (str(tenant_id),))
                            if existing and mode == "create": raise ValueError(f"Academic Terms row {row_no}: term already exists")
                            values=(item["name"],item["term_number"],item["start_date"],item["end_date"],item["is_current"])
                            if existing: await cur.execute("UPDATE academic_terms SET name=%s,term_number=%s,start_date=%s,end_date=%s,is_current=%s WHERE id=%s AND tenant_id=%s",(*values,existing,str(tenant_id))); updated+=1
                            else: await cur.execute("INSERT INTO academic_terms (id,tenant_id,academic_year_id,name,term_number,start_date,end_date,is_current) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),year_id,*values)); created+=1
                        elif sheet == "Departments":
                            existing=await find("departments","code=%s",(item["code"],))
                            if existing and mode=="create": raise ValueError(f"Departments row {row_no}: code already exists")
                            if existing: await cur.execute("UPDATE departments SET name=%s,description=%s WHERE id=%s AND tenant_id=%s",(item["name"],item["description"],existing,str(tenant_id))); updated+=1
                            else: await cur.execute("INSERT INTO departments (id,tenant_id,code,name,description) VALUES (%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),item["code"],item["name"],item["description"])); created+=1
                        elif sheet == "Class Levels":
                            existing=await find("class_levels","code=%s",(item["code"],))
                            if existing and mode=="create": raise ValueError(f"Class Levels row {row_no}: code already exists")
                            if existing: await cur.execute("UPDATE class_levels SET name=%s,level_order=%s WHERE id=%s AND tenant_id=%s",(item["name"],item["level_order"],existing,str(tenant_id))); updated+=1
                            else: await cur.execute("INSERT INTO class_levels (id,tenant_id,code,name,level_order) VALUES (%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),item["code"],item["name"],item["level_order"])); created+=1
                        elif sheet == "Streams":
                            level_id=await require("class_levels","code=%s",(item["class_level"],),f"Class level '{item['class_level']}'")
                            existing=await find("streams","class_level_id=%s AND code=%s",(level_id,item["code"]))
                            if existing and mode=="create": raise ValueError(f"Streams row {row_no}: stream already exists")
                            if existing: await cur.execute("UPDATE streams SET name=%s,capacity=%s WHERE id=%s AND tenant_id=%s",(item["name"],item["capacity"],existing,str(tenant_id))); updated+=1
                            else: await cur.execute("INSERT INTO streams (id,tenant_id,class_level_id,code,name,capacity) VALUES (%s,%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),level_id,item["code"],item["name"],item["capacity"])); created+=1
                        elif sheet == "Subjects":
                            department_id=None
                            if item["department"]: department_id=await require("departments","code=%s",(item["department"],),f"Department '{item['department']}'")
                            existing=await find("subjects","code=%s",(item["code"],))
                            if existing and mode=="create": raise ValueError(f"Subjects row {row_no}: code already exists")
                            if existing: await cur.execute("UPDATE subjects SET department_id=%s,name=%s,short_name=%s,subject_type=%s WHERE id=%s AND tenant_id=%s",(department_id,item["name"],item["short_name"],item["subject_type"],existing,str(tenant_id))); updated+=1
                            else: await cur.execute("INSERT INTO subjects (id,tenant_id,department_id,code,name,short_name,subject_type) VALUES (%s,%s,%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),department_id,item["code"],item["name"],item["short_name"],item["subject_type"])); created+=1
                        elif sheet == "Class Subjects":
                            level_id=await require("class_levels","code=%s",(item["class_level"],),f"Class level '{item['class_level']}'")
                            subject_id=await require("subjects","code=%s",(item["subject"],),f"Subject '{item['subject']}'")
                            existing=await find("class_subjects","class_level_id=%s AND subject_id=%s",(level_id,subject_id))
                            if existing and mode=="create": raise ValueError(f"Class Subjects row {row_no}: assignment already exists")
                            if existing: await cur.execute("UPDATE class_subjects SET is_compulsory=%s WHERE id=%s AND tenant_id=%s",(item["is_compulsory"],existing,str(tenant_id))); updated+=1
                            else: await cur.execute("INSERT INTO class_subjects (id,tenant_id,class_level_id,subject_id,is_compulsory) VALUES (%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),level_id,subject_id,item["is_compulsory"])); created+=1
                        elif sheet == "School Settings":
                            existing=await find("school_settings","setting_key=%s",(item["setting_key"],))
                            if existing and mode=="create": raise ValueError(f"School Settings row {row_no}: key already exists")
                            if existing: await cur.execute("UPDATE school_settings SET setting_value=%s,value_type=%s WHERE id=%s AND tenant_id=%s",(item["setting_value"],item["value_type"],existing,str(tenant_id))); updated+=1
                            else: await cur.execute("INSERT INTO school_settings (id,tenant_id,setting_key,setting_value,value_type) VALUES (%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),item["setting_key"],item["setting_value"],item["value_type"])); created+=1
            await conn.commit()
        except (ValueError, Exception) as exc:
            await conn.rollback()
            if isinstance(exc, ValueError):
                return {"valid": False, "mode": mode, "errors": [{"sheet": "import", "row": None, "field": "data", "message": str(exc)}], "error_count": 1, "created": 0, "updated": 0}
            raise
    return {"valid": True, "mode": mode, "errors": [], "error_count": 0, "created": created, "updated": updated}

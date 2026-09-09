from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from uuid import UUID, uuid4

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.database import get_pool

MAX_ROWS = 5000
MAX_ERRORS = 200
SHEETS = {
    "Students": ["admission_number", "first_name", "middle_name", "last_name", "date_of_birth", "gender", "nationality", "birth_certificate_number", "admission_date", "previous_school", "status", "medical_notes", "emergency_notes"],
    "Guardians": ["first_name", "last_name", "phone", "alternative_phone", "email", "address", "occupation", "employer", "preferred_contact_method", "status"],
    "Student Guardians": ["student_admission_number", "guardian_phone", "guardian_email", "relationship", "is_primary", "is_emergency_contact", "can_pick_up"],
    "Enrollments": ["student_admission_number", "academic_year", "class_level", "stream", "enrollment_date", "exit_date", "status", "notes"],
}


def _text(v, required=False, max_length=None):
    if v is None:
        if required: raise ValueError("value is required")
        return None
    v = str(v).strip()
    if required and not v: raise ValueError("value is required")
    if max_length and len(v) > max_length: raise ValueError(f"maximum length is {max_length}")
    return v or None


def _bool(v, default=False):
    if v is None or v == "": return default
    if isinstance(v, bool): return v
    s = str(v).strip().lower()
    if s in {"true", "1", "yes", "y"}: return True
    if s in {"false", "0", "no", "n"}: return False
    raise ValueError("must be TRUE or FALSE")


def _date(v, required=False):
    if v is None or v == "":
        if required: raise ValueError("date is required")
        return None
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    try: return date.fromisoformat(str(v).strip())
    except ValueError: raise ValueError("must be YYYY-MM-DD")


def _headers(ws, expected):
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return []
    headers = [str(x or "").strip().lower().replace(" (required)", "") for x in rows[0]]
    if headers != expected: raise ValueError("Invalid columns")
    if len(rows) - 1 > MAX_ROWS: raise ValueError(f"Maximum {MAX_ROWS} data rows allowed")
    return [(n, {expected[i]: row[i] if i < len(row) else None for i in range(len(expected))}) for n, row in enumerate(rows[1:], 2) if any(v is not None and str(v).strip() for v in row)]


def _parse(sheet, row_no, row):
    if sheet == "Students":
        return {"admission_number": _text(row["admission_number"], True, 80), "first_name": _text(row["first_name"], True, 100), "middle_name": _text(row["middle_name"], max_length=100), "last_name": _text(row["last_name"], True, 100), "date_of_birth": _date(row["date_of_birth"]), "gender": _text(row["gender"]) or "unspecified", "nationality": _text(row["nationality"], max_length=80), "birth_certificate_number": _text(row["birth_certificate_number"], max_length=100), "admission_date": _date(row["admission_date"]), "previous_school": _text(row["previous_school"], max_length=190), "status": _text(row["status"]) or "active", "medical_notes": _text(row["medical_notes"]), "emergency_notes": _text(row["emergency_notes"])}
    if sheet == "Guardians":
        return {"first_name": _text(row["first_name"], True, 100), "last_name": _text(row["last_name"], True, 100), "phone": _text(row["phone"], max_length=40), "alternative_phone": _text(row["alternative_phone"], max_length=40), "email": _text(row["email"], max_length=190), "address": _text(row["address"], max_length=255), "occupation": _text(row["occupation"], max_length=150), "employer": _text(row["employer"], max_length=190), "preferred_contact_method": _text(row["preferred_contact_method"]) or "phone", "status": _text(row["status"]) or "active"}
    if sheet == "Student Guardians":
        return {"student_admission_number": _text(row["student_admission_number"], True, 80), "guardian_phone": _text(row["guardian_phone"], max_length=40), "guardian_email": _text(row["guardian_email"], max_length=190), "relationship": _text(row["relationship"], True, 80), "is_primary": _bool(row["is_primary"]), "is_emergency_contact": _bool(row["is_emergency_contact"]), "can_pick_up": _bool(row["can_pick_up"], True)}
    if sheet == "Enrollments":
        start, end = _date(row["enrollment_date"], True), _date(row["exit_date"])
        if end and end < start: raise ValueError("exit_date cannot be before enrollment_date")
        return {"student_admission_number": _text(row["student_admission_number"], True, 80), "academic_year": _text(row["academic_year"], True, 190), "class_level": _text(row["class_level"], True, 80), "stream": _text(row["stream"], max_length=80), "enrollment_date": start, "exit_date": end, "status": _text(row["status"]) or "active", "notes": _text(row["notes"], max_length=500)}
    raise ValueError("unsupported sheet")


def _style(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(max(len(str(c.value or "")) for c in col) + 2, 14), 38)


def build_template():
    wb = Workbook(); wb.remove(wb.active)
    info = wb.create_sheet("START HERE")
    info.append(["ShuleLink Students & Guardians Import", ""])
    info.append(["Import order", "Students → Guardians → Student Guardians → Enrollments"])
    info.append(["Rules", "Admission number is the student key. Guardian links use phone or email. Enrollment references Academic Year/Class Level/Stream by their names/codes."])
    info.append(["Modes", "Create New rejects duplicates. Upsert updates matching students/guardians and creates missing records. No data is deleted."])
    info.column_dimensions["A"].width = 28; info.column_dimensions["B"].width = 100
    for sheet, fields in SHEETS.items():
        ws = wb.create_sheet(sheet); ws.append(fields); _style(ws)
    out = BytesIO(); wb.save(out); out.seek(0); return out


async def _export_rows(tenant_id: UUID):
    pool = get_pool(); data = {}
    queries = {
        "Students": "SELECT admission_number,first_name,middle_name,last_name,date_of_birth,gender,nationality,birth_certificate_number,admission_date,previous_school,status,medical_notes,emergency_notes FROM students WHERE tenant_id=%s ORDER BY last_name,first_name,admission_number",
        "Guardians": "SELECT first_name,last_name,phone,alternative_phone,email,address,occupation,employer,preferred_contact_method,status FROM guardians WHERE tenant_id=%s ORDER BY last_name,first_name",
        "Student Guardians": "SELECT s.admission_number,g.phone,g.email,sg.relationship,sg.is_primary,sg.is_emergency_contact,sg.can_pick_up FROM student_guardians sg JOIN students s ON s.id=sg.student_id JOIN guardians g ON g.id=sg.guardian_id WHERE sg.tenant_id=%s ORDER BY s.admission_number,g.last_name",
        "Enrollments": "SELECT s.admission_number,ay.name,cl.code,st.code,e.enrollment_date,e.exit_date,e.status,e.notes FROM student_enrollments e JOIN students s ON s.id=e.student_id JOIN academic_years ay ON ay.id=e.academic_year_id JOIN class_levels cl ON cl.id=e.class_level_id LEFT JOIN streams st ON st.id=e.stream_id WHERE e.tenant_id=%s ORDER BY s.admission_number,e.enrollment_date",
    }
    async with pool.acquire() as conn:
        async with conn.cursor() as cur:
            for sheet, sql in queries.items():
                await cur.execute(sql, (str(tenant_id),)); data[sheet] = await cur.fetchall()
    return data


def build_export(data):
    wb = Workbook(); wb.remove(wb.active)
    info = wb.create_sheet("START HERE"); info.append(["ShuleLink Students & Guardians Export", ""]); info.append(["Import-ready", "Uses admission numbers and stable school-level references rather than database UUIDs."])
    info.column_dimensions["A"].width = 28; info.column_dimensions["B"].width = 100
    for sheet, fields in SHEETS.items():
        ws = wb.create_sheet(sheet); ws.append(fields)
        for row in data.get(sheet, []): ws.append(list(row))
        _style(ws)
        for r in ws.iter_rows():
            for c in r:
                if c.column and "date" in str(ws.cell(1, c.column).value).lower(): c.number_format = "yyyy-mm-dd"
    out = BytesIO(); wb.save(out); out.seek(0); return out


async def import_workbook(tenant_id: UUID, content: bytes, mode: str):
    if len(content) > 10 * 1024 * 1024: raise HTTPException(413, "Import file is too large. Maximum size is 10 MB.")
    try: wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception: raise HTTPException(400, "The uploaded file is not a valid Excel workbook.")
    errors=[]; parsed={}
    try:
        for sheet, fields in SHEETS.items():
            if sheet not in wb.sheetnames: errors.append({"sheet":sheet,"row":1,"message":"Required sheet is missing"}); continue
            try: parsed[sheet]=_headers(wb[sheet], fields)
            except ValueError as exc: errors.append({"sheet":sheet,"row":1,"message":str(exc)})
        for sheet, rows in parsed.items():
            for row_no, raw in rows:
                try: _parse(sheet,row_no,raw)
                except ValueError as exc: errors.append({"sheet":sheet,"row":row_no,"message":str(exc)})
                if len(errors)>=MAX_ERRORS: break
            if len(errors)>=MAX_ERRORS: break
    finally: wb.close()
    if errors: return {"valid":False,"mode":mode,"errors":errors[:MAX_ERRORS],"error_count":len(errors),"created":0,"updated":0}

    pool=get_pool(); created=updated=0
    async with pool.acquire() as conn:
        await conn.begin()
        try:
            async with conn.cursor() as cur:
                students={}; guardians={}; years={}; classes={}; streams={}
                async def one(sql,args):
                    await cur.execute(sql,args); return await cur.fetchone()
                async def student_id(adm):
                    r=await one("SELECT id FROM students WHERE tenant_id=%s AND admission_number=%s",(str(tenant_id),adm)); return str(r[0]) if r else None
                async def guardian_id(phone,email):
                    if phone:
                        r=await one("SELECT id FROM guardians WHERE tenant_id=%s AND phone=%s",(str(tenant_id),phone));
                        if r:return str(r[0])
                    if email:
                        r=await one("SELECT id FROM guardians WHERE tenant_id=%s AND email=%s",(str(tenant_id),email));
                        if r:return str(r[0])
                    return None
                for row_no, raw in parsed.get("Students",[]):
                    d=_parse("Students",row_no,raw); sid=await student_id(d["admission_number"])
                    fields=["first_name","middle_name","last_name","date_of_birth","gender","nationality","birth_certificate_number","admission_date","previous_school","status","medical_notes","emergency_notes"]
                    vals=[d[f] for f in fields]
                    if sid:
                        if mode=="create": raise ValueError(f"Student {d['admission_number']} already exists")
                        await cur.execute(f"UPDATE students SET {','.join(f'{f}=%s' for f in fields)} WHERE id=%s AND tenant_id=%s",(*vals,sid,str(tenant_id))); updated+=1
                    else:
                        sid=str(uuid4()); await cur.execute(f"INSERT INTO students (id,tenant_id,admission_number,{','.join(fields)}) VALUES ({','.join(['%s']*(3+len(fields)))})",(sid,str(tenant_id),d["admission_number"],*vals)); created+=1
                    students[d["admission_number"]]=sid
                for row_no, raw in parsed.get("Guardians",[]):
                    d=_parse("Guardians",row_no,raw); gid=await guardian_id(d["phone"],d["email"])
                    fields=list(d); vals=[d[f] for f in fields]
                    if gid:
                        if mode=="create": raise ValueError("Guardian with the same phone/email already exists")
                        await cur.execute(f"UPDATE guardians SET {','.join(f'{f}=%s' for f in fields)} WHERE id=%s AND tenant_id=%s",(*vals,gid,str(tenant_id))); updated+=1
                    else:
                        gid=str(uuid4()); await cur.execute(f"INSERT INTO guardians (id,tenant_id,{','.join(fields)}) VALUES ({','.join(['%s']*(2+len(fields)))})",(gid,str(tenant_id),*vals)); created+=1
                    guardians[(d["phone"] or "",d["email"] or "")]=gid
                for row_no, raw in parsed.get("Student Guardians",[]):
                    d=_parse("Student Guardians",row_no,raw); sid=students.get(d["student_admission_number"]) or await student_id(d["student_admission_number"])
                    if not sid: raise ValueError(f"Student {d['student_admission_number']} not found")
                    gid=guardians.get((d["guardian_phone"] or "",d["guardian_email"] or "")) or await guardian_id(d["guardian_phone"],d["guardian_email"])
                    if not gid: raise ValueError("Guardian not found by phone/email")
                    if d["is_primary"]: await cur.execute("UPDATE student_guardians SET is_primary=0 WHERE tenant_id=%s AND student_id=%s",(str(tenant_id),sid))
                    existing=await one("SELECT id FROM student_guardians WHERE tenant_id=%s AND student_id=%s AND guardian_id=%s",(str(tenant_id),sid,gid))
                    if existing:
                        if mode=="create": raise ValueError("Student guardian link already exists")
                        await cur.execute("UPDATE student_guardians SET relationship=%s,is_primary=%s,is_emergency_contact=%s,can_pick_up=%s WHERE id=%s",(d["relationship"],d["is_primary"],d["is_emergency_contact"],d["can_pick_up"],str(existing[0]))); updated+=1
                    else:
                        await cur.execute("INSERT INTO student_guardians (id,tenant_id,student_id,guardian_id,relationship,is_primary,is_emergency_contact,can_pick_up) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),sid,gid,d["relationship"],d["is_primary"],d["is_emergency_contact"],d["can_pick_up"])); created+=1
                for row_no, raw in parsed.get("Enrollments",[]):
                    d=_parse("Enrollments",row_no,raw); sid=students.get(d["student_admission_number"]) or await student_id(d["student_admission_number"])
                    if not sid: raise ValueError(f"Student {d['student_admission_number']} not found")
                    ay=await one("SELECT id FROM academic_years WHERE tenant_id=%s AND name=%s",(str(tenant_id),d["academic_year"])); cl=await one("SELECT id FROM class_levels WHERE tenant_id=%s AND code=%s",(str(tenant_id),d["class_level"]))
                    if not ay: raise ValueError(f"Academic year {d['academic_year']} not found")
                    if not cl: raise ValueError(f"Class level {d['class_level']} not found")
                    stream_id=None
                    if d["stream"]:
                        st=await one("SELECT id FROM streams WHERE tenant_id=%s AND class_level_id=%s AND code=%s",(str(tenant_id),str(cl[0]),d["stream"]))
                        if not st: raise ValueError(f"Stream {d['stream']} not found for class {d['class_level']}")
                        stream_id=str(st[0])
                    existing=await one("SELECT id FROM student_enrollments WHERE tenant_id=%s AND student_id=%s AND academic_year_id=%s",(str(tenant_id),sid,str(ay[0])))
                    vals=(str(ay[0]),str(cl[0]),stream_id,d["enrollment_date"],d["exit_date"],d["status"],d["notes"])
                    if existing:
                        if mode=="create": raise ValueError("Student is already enrolled for this academic year")
                        await cur.execute("UPDATE student_enrollments SET class_level_id=%s,stream_id=%s,enrollment_date=%s,exit_date=%s,status=%s,notes=%s WHERE id=%s AND tenant_id=%s",(vals[1],vals[2],vals[3],vals[4],vals[5],vals[6],str(existing[0]),str(tenant_id))); updated+=1
                    else:
                        await cur.execute("INSERT INTO student_enrollments (id,tenant_id,student_id,academic_year_id,class_level_id,stream_id,enrollment_date,exit_date,status,notes) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",(str(uuid4()),str(tenant_id),sid,*vals)); created+=1
            await conn.commit()
        except ValueError:
            await conn.rollback(); raise
        except Exception:
            await conn.rollback(); raise
    return {"valid":True,"mode":mode,"errors":[],"error_count":0,"created":created,"updated":updated}

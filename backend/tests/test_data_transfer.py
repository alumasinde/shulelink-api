from io import BytesIO

from openpyxl import load_workbook

from app.modules.data_transfer.school_structure import DISPLAY_HEADERS, SHEETS, _parse_rows, build_template


def test_school_structure_template_contains_all_import_sheets():
    workbook = load_workbook(build_template(), read_only=True, data_only=True)
    assert workbook.sheetnames[0] == "START HERE"
    assert workbook.sheetnames[1:] == list(SHEETS.keys())
    workbook.close()


def test_school_structure_template_headers_are_stable():
    workbook = load_workbook(build_template(), read_only=True, data_only=True)
    for sheet, fields in SHEETS.items():
        headers = [str(value).strip().lower().replace(" (required)", "") for value in next(workbook[sheet].iter_rows(values_only=True))]
        assert headers == fields
        assert [DISPLAY_HEADERS[field] for field in fields]
    workbook.close()


def test_parser_accepts_empty_valid_workbook():
    workbook = load_workbook(build_template(), read_only=True, data_only=True)
    parsed, errors = _parse_rows(workbook)
    workbook.close()
    assert errors == []
    assert set(parsed) == set(SHEETS)
    assert all(rows == [] for rows in parsed.values())

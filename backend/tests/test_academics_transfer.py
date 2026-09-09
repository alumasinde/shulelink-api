from openpyxl import load_workbook

from app.modules.data_transfer.academics import DISPLAY_HEADERS, SHEETS, _parse, _parse_subjects, _parse_teacher, build_template


def test_teacher_transfer_template_contains_expected_sheets():
    workbook = load_workbook(build_template(), read_only=True, data_only=True)
    assert workbook.sheetnames == ["START HERE", *SHEETS.keys()]
    workbook.close()


def test_teacher_transfer_template_headers_are_stable():
    workbook = load_workbook(build_template(), read_only=True, data_only=True)
    for sheet, fields in SHEETS.items():
        headers = [str(v).strip().lower().replace(" (required)", "") for v in next(workbook[sheet].iter_rows(values_only=True))]
        assert headers == fields
        assert all(field in DISPLAY_HEADERS for field in fields)
    workbook.close()


def test_teacher_parser_rejects_invalid_status():
    errors = []
    row = {
        "teacher_number": "T001", "first_name": "Jane", "middle_name": None,
        "last_name": "Doe", "gender": "female", "phone": None, "email": None,
        "department": "Languages", "employment_type": None, "employment_date": None,
        "status": "retired", "notes": None,
    }
    assert _parse_teacher(2, row, errors) is None
    assert "status must be active, inactive, on_leave, or terminated" in errors[0]["message"]


def test_subject_parser_limits_duplicate_subjects():
    errors = []
    row = {"teacher_number": "T001", "subject_1": "ENG", "subject_2": " eng "}
    assert _parse_subjects(2, row, errors) is None
    assert "same subject" in errors[0]["message"]


def test_parser_accepts_template():
    workbook = load_workbook(build_template(), read_only=True, data_only=True)
    parsed, errors = _parse(workbook)
    workbook.close()
    assert errors == []
    assert set(parsed) == set(SHEETS)
    assert all(rows == [] for rows in parsed.values())
